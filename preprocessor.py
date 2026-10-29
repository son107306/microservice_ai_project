
"""Module trích xuất văn bản từ PDF, DOCX và PPTX."""

import os
import unicodedata
from io import BytesIO

import pymupdf as fitz
import docx
from pptx import Presentation


SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".pptx"}


def clean_and_normalize_text(text: str) -> str:
    """Chuẩn hóa Unicode và loại bỏ dòng trống thừa."""
    if not text:
        return ""

    normalized = unicodedata.normalize("NFKC", text)
    lines = [
        line.strip()
        for line in normalized.splitlines()
        if line.strip()
    ]
    return "\n".join(lines)


def extract_from_pdf(
    pdf_bytes: bytes,
    is_4up_slides: bool = False,
) -> str:
    """Trích xuất văn bản từ PDF thường hoặc PDF handout 4-up."""
    if not pdf_bytes:
        raise ValueError("File PDF rỗng.")

    output_blocks = []

    try:
        with fitz.open(stream=pdf_bytes, filetype="pdf") as pdf:
            for page_idx, page in enumerate(pdf, start=1):
                if is_4up_slides:
                    width = page.rect.width
                    height = page.rect.height

                    quadrants = [
                        (
                            "Trên-Trái",
                            fitz.Rect(0, 0, width / 2, height / 2),
                        ),
                        (
                            "Trên-Phải",
                            fitz.Rect(width / 2, 0, width, height / 2),
                        ),
                        (
                            "Dưới-Trái",
                            fitz.Rect(0, height / 2, width / 2, height),
                        ),
                        (
                            "Dưới-Phải",
                            fitz.Rect(
                                width / 2,
                                height / 2,
                                width,
                                height,
                            ),
                        ),
                    ]

                    for label, rect in quadrants:
                        text = page.get_text("text", clip=rect).strip()

                        if text:
                            output_blocks.append(
                                f"=== Trang {page_idx} - {label} ===\n{text}"
                            )
                else:
                    text = page.get_text("text").strip()

                    if text:
                        output_blocks.append(
                            f"=== Trang {page_idx} ===\n{text}"
                        )

    except Exception as exc:
        raise ValueError(f"Không thể đọc PDF: {exc}") from exc

    return clean_and_normalize_text("\n\n".join(output_blocks))


def extract_from_docx(docx_bytes: bytes) -> str:
    """Trích xuất đoạn văn và bảng từ file Word DOCX."""
    if not docx_bytes:
        raise ValueError("File Word rỗng.")

    try:
        document = docx.Document(BytesIO(docx_bytes))
    except Exception as exc:
        raise ValueError(f"Không thể đọc DOCX: {exc}") from exc

    output = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()

        if text:
            output.append(text)

    for table_idx, table in enumerate(document.tables, start=1):
        output.append(f"[Bảng {table_idx}]")

        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]

            if any(cells):
                output.append(" | ".join(cells))

    return clean_and_normalize_text("\n".join(output))


def extract_from_pptx(pptx_bytes: bytes) -> str:
    """Trích xuất văn bản, bảng và ghi chú từ PowerPoint PPTX."""
    if not pptx_bytes:
        raise ValueError("File PowerPoint rỗng.")

    try:
        presentation = Presentation(BytesIO(pptx_bytes))
    except Exception as exc:
        raise ValueError(f"Không thể đọc PPTX: {exc}") from exc

    output_slides = []

    for slide_idx, slide in enumerate(presentation.slides, start=1):
        slide_parts = []

        for shape in slide.shapes:
            if shape.has_text_frame:
                text = shape.text.strip()

                if text:
                    slide_parts.append(text)

            if shape.has_table:
                for row in shape.table.rows:
                    cells = [cell.text.strip() for cell in row.cells]

                    if any(cells):
                        slide_parts.append(" | ".join(cells))

        if slide.has_notes_slide:
            notes_frame = slide.notes_slide.notes_text_frame

            if notes_frame is not None:
                notes = "\n".join(
                    paragraph.text.strip()
                    for paragraph in notes_frame.paragraphs
                    if paragraph.text.strip()
                    and not paragraph.text.strip().isdigit()
                )

                if notes:
                    slide_parts.append(f"[Ghi chú: {notes}]")

        if slide_parts:
            output_slides.append(
                f"=== Slide {slide_idx} ===\n"
                + "\n".join(slide_parts)
            )

    return clean_and_normalize_text("\n\n".join(output_slides))


def extract_document(
    file_bytes: bytes,
    filename: str,
    is_4up_slides: bool = False,
) -> str:
    """Tự chọn hàm trích xuất dựa trên phần mở rộng của file."""
    if not file_bytes:
        raise ValueError("File tải lên rỗng.")

    ext = os.path.splitext(filename)[1].lower()

    if ext not in SUPPORTED_EXTENSIONS:
        raise ValueError(
            f"Định dạng {ext or '(không có phần mở rộng)'} "
            "chưa được hỗ trợ. Vui lòng sử dụng PDF, DOCX hoặc PPTX."
        )

    if ext == ".pdf":
        result = extract_from_pdf(
            file_bytes,
            is_4up_slides=is_4up_slides,
        )
    elif ext == ".docx":
        result = extract_from_docx(file_bytes)
    else:
        result = extract_from_pptx(file_bytes)

    if not result.strip():
        raise ValueError(
            "Không trích xuất được văn bản. "
            "File có thể chỉ chứa ảnh hoặc nội dung không được hỗ trợ."
        )

    return result


if __name__ == "__main__":
    sample = (
        "Công thức: ∑ Fx = 0; "
        "góc α = 30°; độ lệch Δ = 0.05"
    )

    print("=== KIỂM THỬ CHUẨN HÓA VĂN BẢN ===")
    print(clean_and_normalize_text(sample))

    print("\nCác hàm trích xuất đã được khai báo:")
    print("- extract_from_pdf")
    print("- extract_from_docx")
    print("- extract_from_pptx")
    print("- extract_document")
    print("\nLưu ý: cần kiểm thử bằng file mẫu thật.")
