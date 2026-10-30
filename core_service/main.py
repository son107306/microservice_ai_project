
from uuid import uuid4
from datetime import datetime, timezone
from typing import Any
from io import BytesIO

import pymupdf as fitz
from docx import Document
from pptx import Presentation

from fastapi import (
    FastAPI,
    HTTPException,
    UploadFile,
    File,
    Form,
)
from pydantic import BaseModel, Field

from preprocessor import extract_document


app = FastAPI(
    title="Quiz Generator - Core Service",
    description="Quản lý tài liệu và kết quả trắc nghiệm",
    version="1.0.0",
)


# Lưu trữ dữ liệu tạm thời trong RAM
documents: dict[str, dict[str, Any]] = {}
quizzes: dict[str, dict[str, Any]] = {}


# =========================
# PYDANTIC MODELS
# =========================

class DocumentCreate(BaseModel):
    filename: str = Field(min_length=1, max_length=255)
    page_count: int = Field(ge=1)
    character_count: int = Field(ge=0)
    content: str = Field(min_length=1)


class QuizCreate(BaseModel):
    document_id: str
    questions: list[dict[str, Any]]


# =========================
# HELPER FUNCTIONS
# =========================

def get_page_count(file_bytes: bytes, filename: str) -> int:
    extension = filename.lower().split(".")[-1]

    if extension == "pdf":
        with fitz.open(
            stream=file_bytes,
            filetype="pdf"
        ) as doc:
            return len(doc)

    if extension == "docx":
        doc = Document(BytesIO(file_bytes))
        return max(1, len(doc.paragraphs) // 30 + 1)

    if extension == "pptx":
        presentation = Presentation(BytesIO(file_bytes))
        return len(presentation.slides)

    return 1


# =========================
# HEALTH CHECK
# =========================

@app.get("/health")
def health():
    return {
        "status": "ok",
        "service": "core-service"
    }


# =========================
# DOCUMENT APIs
# =========================

# Tạo thông tin tài liệu từ dữ liệu JSON
@app.post("/api/v1/documents", status_code=201)
def create_document(data: DocumentCreate):
    document_id = str(uuid4())

    document = {
        "document_id": document_id,
        **data.model_dump(),
        "created_at": datetime.now(timezone.utc).isoformat(),
    }

    documents[document_id] = document

    return document


# Upload và trích xuất nội dung tài liệu
@app.post("/api/v1/documents/upload", status_code=201)
async def upload_document(
    file: UploadFile = File(...),
    is_4up_slides: bool = Form(False)
):
    filename = file.filename or ""

    allowed_extensions = (".pdf", ".docx", ".pptx")

    if not filename.lower().endswith(allowed_extensions):
        raise HTTPException(
            status_code=400,
            detail="Only PDF, DOCX and PPTX files are supported"
        )

    file_bytes = await file.read()

    if not file_bytes:
        raise HTTPException(
            status_code=400,
            detail="Uploaded file is empty"
        )

    try:
        # Trích xuất văn bản từ tài liệu
        content = extract_document(
            file_bytes=file_bytes,
            filename=filename,
            is_4up_slides=is_4up_slides
        )

        if not content or not content.strip():
            raise HTTPException(
                status_code=400,
                detail="No text could be extracted from the document"
            )

        # Đếm số trang hoặc slide
        page_count = get_page_count(
            file_bytes,
            filename
        )

        # Tạo dữ liệu tài liệu
        data = DocumentCreate(
            filename=filename,
            page_count=page_count,
            character_count=len(content),
            content=content
        )

        # Lưu tài liệu
        return create_document(data)

    except HTTPException:
        raise

    except Exception as exc:
        raise HTTPException(
            status_code=400,
            detail=f"Document processing failed: {str(exc)}"
        )


# Lấy thông tin tài liệu
@app.get("/api/v1/documents/{document_id}")
def get_document(document_id: str):
    document = documents.get(document_id)

    if document is None:
        raise HTTPException(
            status_code=404,
            detail="Document not found"
        )

    return document


# =========================
# QUIZ APIs
# =========================

# Tạo quiz
@app.post("/api/v1/quizzes", status_code=201)
def create_quiz(data: QuizCreate):
    if data.document_id not in documents:
        raise HTTPException(
            status_code=404,
            detail="Document not found"
        )

    quiz_id = str(uuid4())

    quiz = {
        "quiz_id": quiz_id,
        "document_id": data.document_id,
        "questions": data.questions,
        "question_count": len(data.questions),
        "created_at": datetime.now(timezone.utc).isoformat(),
    }

    quizzes[quiz_id] = quiz

    return quiz


# Lấy kết quả quiz
@app.get("/api/v1/quizzes/{quiz_id}")
def get_quiz(quiz_id: str):
    quiz = quizzes.get(quiz_id)

    if quiz is None:
        raise HTTPException(
            status_code=404,
            detail="Quiz not found"
        )

    return quiz
